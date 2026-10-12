import json
from pathlib import Path
from uuid import UUID
import asyncio

from docx import Document as DocxDocument
from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    HTTPException,
    UploadFile,
)
from fastapi.responses import StreamingResponse

from pypdf import PdfReader
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance,
    PointStruct,
    VectorParams,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from auth import get_current_user
from config import settings
from database import SessionLocal, get_db
from embeddings import embed_text
from models import (
    Document,
    DocumentChunk,
    KnowledgeBase,
    User,
)
from worker import process_document_task


# ============================================================
# ROUTER
# ============================================================

router = APIRouter(
    prefix="/knowledge-bases",
    tags=["Documents"],
)


# ============================================================
# UPLOAD CONFIGURATION
# ============================================================

UPLOAD_DIR = Path("uploads")

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".docx",
    ".json",
}


# ============================================================
# QDRANT
# ============================================================

qdrant_client = QdrantClient(
    url=settings.qdrant_url,
    api_key=settings.qdrant_api_key,
)


_named_vectors_cache: bool | None = None


def get_vector_payload(vector: list[float]):
    global _named_vectors_cache
    if _named_vectors_cache is None:
        try:
            col = qdrant_client.get_collection(settings.qdrant_collection_name)
            _named_vectors_cache = isinstance(col.config.params.vectors, dict) and "dense" in col.config.params.vectors
        except Exception:
            _named_vectors_cache = True
    return {"dense": vector} if _named_vectors_cache else vector


def ensure_qdrant_collection():

    collections = (
        qdrant_client
        .get_collections()
        .collections
    )

    collection_names = {
        collection.name
        for collection in collections
    }

    if settings.qdrant_collection_name not in collection_names:

        qdrant_client.create_collection(

            collection_name=
                settings.qdrant_collection_name,

            vectors_config={
                "dense": VectorParams(
                    size=settings.dense_embedding_dim,
                    distance=Distance.COSINE,
                )
            },
        )


# ============================================================
# DOCUMENT LOADER
# ============================================================

class DocumentLoader:
    """
    Loads text from supported document types.
    """

    def load(
        self,
        file_path: str,
        file_type: str,
    ) -> tuple[str, list[dict]]:

        extension = file_type.lower()

        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        if extension == ".txt":

            text = Path(
                file_path
            ).read_text(
                encoding="utf-8"
            )

            return text, [
                {
                    "page": None,
                    "text": text,
                }
            ]

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        if extension == ".pdf":

            reader = PdfReader(
                file_path
            )

            pages = []

            for page_number, page in enumerate(
                reader.pages,
                start=1,
            ):

                page_text = (
                    page.extract_text()
                    or ""
                )

                pages.append(
                    {
                        "page": page_number,
                        "text": page_text,
                    }
                )

            full_text = "\n\n".join(
                page["text"]
                for page in pages
            )

            return full_text, pages

        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        if extension == ".docx":

            doc = DocxDocument(
                file_path
            )

            paragraphs = [
                paragraph.text
                for paragraph in doc.paragraphs
                if paragraph.text.strip()
            ]

            text = "\n".join(
                paragraphs
            )

            return text, [
                {
                    "page": None,
                    "text": text,
                }
            ]

        # ----------------------------------------------------
        # JSON
        # ----------------------------------------------------

        if extension == ".json":

            with open(
                file_path,
                "r",
                encoding="utf-8",
            ) as file:

                data = json.load(file)

            text = json.dumps(
                data,
                indent=2,
                ensure_ascii=False,
            )

            return text, [
                {
                    "page": None,
                    "text": text,
                }
            ]

        raise ValueError(
            f"Unsupported file type: {file_type}"
        )


# ============================================================
# TEXT CHUNKER
# ============================================================

class TextChunker:

    def __init__(
        self,
        chunk_size: int = 800,
        chunk_overlap: int = 150,
    ):

        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def chunk(
        self,
        text: str,
    ) -> list[str]:

        text = text.strip()

        if not text:
            return []

        chunks = []

        start = 0
        text_length = len(text)

        while start < text_length:

            end = start + self.chunk_size

            chunk = text[start:end].strip()

            if chunk:
                chunks.append(chunk)

            if end >= text_length:
                break

            start = end - self.chunk_overlap

        return chunks


# ============================================================
# KNOWLEDGE BASE CREATION
# ============================================================

@router.post(
    "",
)
async def create_knowledge_base(
    name: str,
    description: str | None = None,
    current_user: User = Depends(
        get_current_user
    ),
    db: AsyncSession = Depends(
        get_db
    ),
):

    knowledge_base = KnowledgeBase(

        user_id=current_user.id,

        name=name,

        description=description,
    )

    db.add(
        knowledge_base
    )

    await db.commit()

    await db.refresh(
        knowledge_base
    )

    return {
        "id": str(
            knowledge_base.id
        ),
        "name":
            knowledge_base.name,
        "description":
            knowledge_base.description,
    }


# ============================================================
# GET KNOWLEDGE BASES
# ============================================================

@router.get("")
async def get_knowledge_bases(
    current_user: User = Depends(
        get_current_user
    ),
    db: AsyncSession = Depends(
        get_db
    ),
):

    result = await db.execute(

        select(KnowledgeBase)
        .where(
            KnowledgeBase.user_id
            == current_user.id
        )
        .order_by(
            KnowledgeBase.created_at.desc()
        )
    )

    knowledge_bases = (
        result.scalars().all()
    )

    return knowledge_bases


# ============================================================
# UPLOAD DOCUMENT
# ============================================================

@router.post(
    "/{knowledge_base_id}/documents"
)
async def upload_document(
    knowledge_base_id: UUID,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    current_user: User = Depends(
        get_current_user
    ),
    db: AsyncSession = Depends(
        get_db
    ),
):

    # --------------------------------------------------------
    # 1. Verify knowledge base ownership
    # --------------------------------------------------------

    result = await db.execute(

        select(KnowledgeBase).where(

            KnowledgeBase.id
            == knowledge_base_id,

            KnowledgeBase.user_id
            == current_user.id,
        )
    )

    knowledge_base = (
        result.scalar_one_or_none()
    )

    if knowledge_base is None:

        raise HTTPException(
            status_code=404,
            detail="Knowledge base not found",
        )

    # --------------------------------------------------------
    # 2. Validate filename
    # --------------------------------------------------------

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="Filename is required",
        )

    extension = (
        Path(file.filename)
        .suffix
        .lower()
    )

    if extension not in ALLOWED_EXTENSIONS:

        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Allowed: PDF, TXT, DOCX, JSON"
            ),
        )

    # --------------------------------------------------------
    # 3. Read file
    # --------------------------------------------------------

    file_data = await file.read()

    max_size = (
        settings.max_upload_size_mb
        * 1024
        * 1024
    )

    if len(file_data) > max_size:

        raise HTTPException(
            status_code=413,
            detail=(
                f"File exceeds "
                f"{settings.max_upload_size_mb} MB"
            ),
        )

    # --------------------------------------------------------
    # 4. Create user-specific directory
    # --------------------------------------------------------

    user_directory = (
        UPLOAD_DIR
        / str(current_user.id)
        / str(knowledge_base_id)
    )

    user_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------------
    # 5. Save file
    # --------------------------------------------------------

    file_path = (
        user_directory
        / file.filename
    )

    file_path.write_bytes(
        file_data
    )

    # --------------------------------------------------------
    # 6. Create document DB record
    # --------------------------------------------------------

    document = Document(

        knowledge_base_id=
            knowledge_base_id,

        filename=
            file.filename,

        file_type=
            extension,

        file_path=
            str(file_path),

        status=
            "pending",
    )

    db.add(
        document
    )

    await db.commit()

    await db.refresh(
        document
    )

    # --------------------------------------------------------
    # 7. Queue background task (Celery with native fallback)
    # --------------------------------------------------------

    try:
        process_document_task.delay(
            str(document.id)
        )
    except Exception as exc:
        print(
            f"[INGESTION] Celery dispatch skipped ({exc}), "
            f"processing with BackgroundTasks"
        )
        background_tasks.add_task(
            process_document_by_id,
            str(document.id),
        )

    # --------------------------------------------------------
    # 8. Return immediately
    # --------------------------------------------------------

    return {

        "document_id":
            str(document.id),

        "filename":
            document.filename,

        "status":
            "pending",

        "message":
            "Document uploaded and queued for processing",
    }


# ============================================================
# PROCESS DOCUMENT
# ============================================================

async def process_document_by_id(
    document_id: str,
):

    document = None

    async with SessionLocal() as db:

        try:

            # ------------------------------------------------
            # 1. Find document
            # ------------------------------------------------

            result = await db.execute(

                select(Document).where(

                    Document.id
                    == UUID(document_id)
                )
            )

            document = (
                result.scalar_one_or_none()
            )

            if document is None:

                raise ValueError(
                    f"Document {document_id} "
                    f"not found"
                )

            # ------------------------------------------------
            # 2. Mark processing
            # ------------------------------------------------

            document.status = "processing"

            document.error_message = None

            await db.commit()

            # ------------------------------------------------
            # 3. Load document
            # ------------------------------------------------

            loader = DocumentLoader()

            full_text, pages = (
                loader.load(
                    document.file_path,
                    document.file_type,
                )
            )

            if not full_text.strip():

                raise ValueError(
                    "Document contains no readable text"
                )

            # ------------------------------------------------
            # 4. Chunk document
            # ------------------------------------------------

            chunker = TextChunker(

                chunk_size=
                    settings.chunk_size,

                chunk_overlap=
                    settings.chunk_overlap,
            )

            chunks = chunker.chunk(
                full_text
            )

            if not chunks:

                raise ValueError(
                    "No chunks were created"
                )

            # ------------------------------------------------
            # 5. Process chunks
            # ------------------------------------------------

            ensure_qdrant_collection()

            total_chunks = 0

            for index, chunk_text in enumerate(
                chunks
            ):

                print(
                    f"[WORKER] "
                    f"Processing chunk "
                    f"{index + 1}/"
                    f"{len(chunks)}"
                )

                # --------------------------------------------
                # Embedding
                # --------------------------------------------

                vector = embed_text(
                    chunk_text
                )

                # --------------------------------------------
                # Determine page
                # --------------------------------------------

                page_number = None

                if document.file_type == ".pdf":

                    current_position = (
                        full_text.find(
                            chunk_text
                        )
                    )

                    if current_position >= 0:

                        characters_seen = 0

                        for page in pages:

                            characters_seen += len(
                                page["text"]
                            )

                            if (
                                current_position
                                <= characters_seen
                            ):

                                page_number = (
                                    page["page"]
                                )

                                break

                # --------------------------------------------
                # PostgreSQL chunk
                # --------------------------------------------

                chunk = DocumentChunk(

                    document_id=
                        document.id,

                    chunk_index=
                        index,

                    content=
                        chunk_text,

                    page_number=
                        page_number,

                    token_count=
                        len(
                            chunk_text.split()
                        ),
                )

                db.add(
                    chunk
                )

                await db.flush()

                # --------------------------------------------
                # Qdrant point
                # --------------------------------------------

                point_id = str(
                    chunk.id
                )

                qdrant_client.upsert(

                    collection_name=
                        settings.qdrant_collection_name,

                    points=[
                        PointStruct(

                            id=point_id,

                            vector=get_vector_payload(
                                vector
                            ),

                            payload={

                                "knowledge_base_id":
                                    str(
                                        document.knowledge_base_id
                                    ),

                                "document_id":
                                    str(
                                        document.id
                                    ),

                                "chunk_id":
                                    point_id,

                                "source":
                                    document.filename,

                                "page":
                                    page_number,

                                "chunk_index":
                                    index,

                                "text":
                                    chunk_text,
                            },
                        )
                    ],
                )

                # --------------------------------------------
                # Save Qdrant point ID
                # --------------------------------------------

                chunk.qdrant_point_id = (
                    point_id
                )

                total_chunks += 1

            # ------------------------------------------------
            # 6. Mark processed
            # ------------------------------------------------

            document.chunk_count = (
                total_chunks
            )

            document.status = (
                "processed"
            )

            document.error_message = None

            await db.commit()

            print(
                f"[WORKER] "
                f"Document processed: "
                f"{document.filename}"
            )

            return {

                "document_id":
                    document_id,

                "status":
                    "processed",

                "chunk_count":
                    total_chunks,
            }

        except Exception as error:

            # ------------------------------------------------
            # 7. Mark failed
            # ------------------------------------------------

            if document is not None:

                document.status = "failed"

                document.error_message = str(
                    error
                )

                await db.commit()

            print(
                f"[WORKER] "
                f"Document processing failed: "
                f"{error}"
            )

            raise


# ============================================================
# GET DOCUMENT STATUS
# ============================================================

@router.get(
    "/{knowledge_base_id}/documents"
)
async def get_documents(
    knowledge_base_id: UUID,
    current_user: User = Depends(
        get_current_user
    ),
    db: AsyncSession = Depends(
        get_db
    ),
):

    # --------------------------------------------------------
    # Verify KB ownership
    # --------------------------------------------------------

    result = await db.execute(

        select(KnowledgeBase).where(

            KnowledgeBase.id
            == knowledge_base_id,

            KnowledgeBase.user_id
            == current_user.id,
        )
    )

    knowledge_base = (
        result.scalar_one_or_none()
    )

    if knowledge_base is None:

        raise HTTPException(
            status_code=404,
            detail="Knowledge base not found",
        )

    # --------------------------------------------------------
    # Get documents
    # --------------------------------------------------------

    result = await db.execute(

        select(Document)
        .where(
            Document.knowledge_base_id
            == knowledge_base_id
        )
        .order_by(
            Document.created_at.desc()
        )
    )

    documents = (
        result.scalars().all()
    )

    return [

        {
            "id":
                str(document.id),

            "filename":
                document.filename,

            "file_type":
                document.file_type,

            "status":
                document.status,

            "chunk_count":
                document.chunk_count,

            "error_message":
                document.error_message,

            "created_at":
                document.created_at,
        }

        for document in documents
    ]
@router.get("/documents/{document_id}/status")
async def get_document_status(
    document_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Document)
        .join(KnowledgeBase, Document.knowledge_base_id == KnowledgeBase.id)
        .where(
            Document.id == document_id,
            KnowledgeBase.user_id == current_user.id,
        )
    )

    document = result.scalar_one_or_none()

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found",
        )

    return {
        "document_id": str(document.id),
        "filename": document.filename,
        "status": document.status,
        "chunk_count": document.chunk_count,
        "error_message": document.error_message,
    }

@router.get("/documents/{document_id}/stream")
async def stream_document_status(
    document_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Document)
        .join(
            KnowledgeBase,
            Document.knowledge_base_id == KnowledgeBase.id,
        )
        .where(
            Document.id == document_id,
            KnowledgeBase.user_id == current_user.id,
        )
    )

    document = result.scalar_one_or_none()

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found",
        )

    async def event_generator():
        while True:
            result = await db.execute(
                select(Document).where(
                    Document.id == document_id
                )
            )

            document = result.scalar_one_or_none()

            if document is None:
                break

            data = {
                "document_id": str(document.id),
                "status": document.status,
                "chunk_count": document.chunk_count,
                "error_message": document.error_message,
            }

            yield f"data: {json.dumps(data)}\n\n"

            if document.status in {"processed", "failed"}:
                break

            await asyncio.sleep(2)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        },
    )